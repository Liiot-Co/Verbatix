"""Regresión: parse_verbatims exige Respondent ID (cruce con perfil) y admite
column_map explícito. El .dau no pasa por esta función."""
import io

import pytest
from openpyxl import Workbook

from app.pipeline.ingest import excel


def _make_xlsx(header, data_rows):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in data_rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_rows_returns_columns_and_dict_rows():
    content = _make_xlsx(
        ["Marca", "Razon"],
        [["coca cola", "sabor"], ["pepsi", "precio"]],
    )
    columns, rows = excel.read_rows(content, "Sheet")
    assert columns == ["Marca", "Razon"]
    assert rows == [
        {"Marca": "coca cola", "Razon": "sabor"},
        {"Marca": "pepsi", "Razon": "precio"},
    ]


def test_read_rows_empty_sheet():
    content = _make_xlsx([], [])
    columns, rows = excel.read_rows(content, "Sheet")
    assert columns == []
    assert rows == []


def test_parse_verbatims_without_column_map_unchanged():
    content = _make_xlsx(
        ["Respondent ID", "Question ID", "Verbatim"],
        [["1", "Q1", "sabor rico"], ["2", "Q1", "buen precio"]],
    )
    rows = excel.parse_verbatims(content, "Sheet")
    assert len(rows) == 2
    assert rows[0]["original_text"] == "sabor rico"
    assert rows[0]["respondent_external_id"] == "1"
    assert rows[1]["respondent_external_id"] == "2"


def test_parse_verbatims_with_explicit_column_map():
    content = _make_xlsx(
        ["Respondent ID", "Question ID", "Verbatim", "Comentario libre"],
        [["1", "Q1", "sabor rico", "otro texto"]],
    )
    rows = excel.parse_verbatims(
        content,
        "Sheet",
        {"verbatim": "Comentario libre", "respondent_id": "Respondent ID"},
    )
    assert len(rows) == 1
    assert rows[0]["original_text"] == "otro texto"
    assert rows[0]["respondent_external_id"] == "1"


def test_parse_verbatims_excel_solo_respondent_y_verbatim():
    """Formato típico Excel (sin Question ID): Respondent ID + Verbatim."""
    content = _make_xlsx(
        ["Respondent ID", "Verbatim"],
        [["101", "me gusta"], ["102", "no me gusta"], ["", "sin id"]],
    )
    rows = excel.parse_verbatims(content, "Sheet")
    assert len(rows) == 2
    assert {r["respondent_external_id"] for r in rows} == {"101", "102"}
    assert all(r["question_key"] == "verbatim" for r in rows)


def test_parse_verbatims_accepts_serial_alias():
    content = _make_xlsx(
        ["Serial", "Verbatim"],
        [["7", "texto"]],
    )
    rows = excel.parse_verbatims(content, "Sheet")
    assert rows[0]["respondent_external_id"] == "7"


def test_parse_verbatims_rejects_missing_respondent_id():
    content = _make_xlsx(
        ["Verbatim"],
        [["solo texto"]],
    )
    with pytest.raises(ValueError, match="Respondent ID"):
        excel.parse_verbatims(content, "Sheet")


def test_parse_verbatims_mapped_respondent_column():
    content = _make_xlsx(
        ["ID encuestado", "Respuesta abierta"],
        [["55", "buen sabor"]],
    )
    rows = excel.parse_verbatims(
        content,
        "Sheet",
        {"respondent_id": "ID encuestado", "verbatim": "Respuesta abierta"},
    )
    assert len(rows) == 1
    assert rows[0]["respondent_external_id"] == "55"
    assert rows[0]["original_text"] == "buen sabor"


def test_parse_verbatims_preserves_realid_and_source_row():
    content = _make_xlsx(
        ["Respondent ID", "realid", "Question ID", "Verbatim"],
        [["7", "persona-abc", "Q1", "sabor"]],
    )
    rows = excel.parse_verbatims(content, "Sheet")
    assert rows[0]["respondent_external_id"] == "7"
    assert rows[0]["realid"] == "persona-abc"
    assert rows[0]["source_row_number"] == 2


def test_parse_verbatims_two_sheets_concatenate_without_cross_contamination():
    """El router (`POST .../import`, `extra_sheets`) parsea la hoja principal y
    cada hoja extra por separado y SUMA sus filas antes de un único
    insert_verbatims — esto reproduce esa concatenación a nivel de parser: dos
    hojas separadas (Marcas / Razones), mismo Respondent ID por persona, deben
    combinarse sin que el source_row_number (reiniciado por hoja) ni el
    question_key de una se filtre a la otra."""
    from openpyxl import Workbook

    wb = Workbook()
    marcas_sheet = wb.active
    marcas_sheet.title = "Marcas"
    marcas_sheet.append(["Respondent ID", "Marca mencionada"])
    marcas_sheet.append(["1", "gloria"])
    razones_sheet = wb.create_sheet("Razones")
    razones_sheet.append(["Respondent ID", "Motivo"])
    razones_sheet.append(["1", "buen sabor"])
    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    brand_rows = excel.parse_verbatims(
        content, "Marcas", {"respondent_id": "Respondent ID", "brands": "Marca mencionada"}
    )
    reason_rows = excel.parse_verbatims(
        content, "Razones", {"respondent_id": "Respondent ID", "reasons": "Motivo"}
    )
    combined = brand_rows + reason_rows

    assert len(combined) == 2
    assert {r["question_key"] for r in combined} == {"brands", "reasons"}
    # Mismo Respondent ID en ambas hojas -> mismo respondent_external_id, para
    # que insert_verbatims los cuelgue del mismo encuestado (conflict key
    # (round_id, imported_file_id, respondent_external_id) en `respondents`).
    assert {r["respondent_external_id"] for r in combined} == {"1"}
    assert combined[0]["original_text"] == "gloria"
    assert combined[1]["original_text"] == "buen sabor"
