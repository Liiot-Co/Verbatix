"""Generador del Excel del Libro de Códigos (LC)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _sheet_suffix(coding_type: str) -> str:
    return "Marcas" if coding_type == "brands" else "Razones"


# Estilos que replican el formato del Libro de Códigos de referencia
# (LC Marcas de los estudios reales: título amarillo, encabezado amarillo,
# filas NETO amarillas, código/frecuencia/% en negrita, contenido en wrap).
_HEADER_FILL = PatternFill("solid", fgColor="FFC000")
_TITLE_FONT = Font(name="Verdana", size=12, bold=True)
_HEADER_FONT = Font(name="Verdana", size=12, bold=True)
_NET_FONT = Font(name="Verdana", size=12, bold=True)
_CODE_FONT = Font(name="Verdana", size=12, bold=True)
_LABEL_FONT = Font(name="Verdana", size=11)
_CONTENT_FONT = Font(name="Verdana", size=11)
_METRIC_FONT = Font(name="Verdana", size=12, bold=True, color="FF0000")

_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

_THIN = Side(style="thin")
_MEDIUM = Side(style="medium")

_LC_COLUMN_WIDTHS = {"A": 4.45, "B": 42, "C": 60.73, "D": 54.45, "E": 16.82, "F": 7.18}
_LC_LAST_COLUMN = 6


def _row_border(col: int) -> Border:
    """Borde de una fila normal (no fusionada): caja completa, con 'medium'
    en la derecha cuando es el borde exterior de la tabla (columna F)."""
    right = _MEDIUM if col == _LC_LAST_COLUMN else _THIN
    return Border(top=_THIN, bottom=_THIN, left=_THIN, right=right)


def _merged_row_border(col: int) -> Border:
    """Borde de una fila fusionada (título / NETO): solo fino, sin el acento
    'medium' del borde derecho (ese acento solo aplica a filas normales)."""
    if col == 1:
        return Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
    if col == _LC_LAST_COLUMN:
        return Border(top=_THIN, bottom=_THIN, left=None, right=_THIN)
    return Border(top=_THIN, bottom=_THIN, left=None, right=None)


def _style_lc_sheet(lc_ws, title: str) -> None:
    for letter, width in _LC_COLUMN_WIDTHS.items():
        lc_ws.column_dimensions[letter].width = width
    lc_ws.sheet_view.showGridLines = False
    lc_ws.sheet_view.zoomScale = 80

    lc_ws.row_dimensions[1].height = 38.15
    lc_ws.row_dimensions[2].height = 15

    lc_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=_LC_LAST_COLUMN)
    lc_ws.cell(row=1, column=1, value=title)
    for col in range(1, _LC_LAST_COLUMN + 1):
        cell = lc_ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _TITLE_FONT
        cell.alignment = _CENTER_WRAP
        cell.border = _merged_row_border(col)

    headers = ["", "CÓDIGO", "MENCIÓN PRINCIPAL", "CONTENIDO", "FRECUENCIA", "%"]
    for col, header in enumerate(headers, start=1):
        cell = lc_ws.cell(row=2, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _row_border(col)


def _write_lc_sheet(lc_ws, nets: list[dict]) -> None:
    row = 3
    for net in nets or []:
        lc_ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=_LC_LAST_COLUMN)
        lc_ws.cell(row=row, column=1, value=net["name"])
        for col in range(1, _LC_LAST_COLUMN + 1):
            cell = lc_ws.cell(row=row, column=col)
            cell.fill = _HEADER_FILL
            cell.font = _NET_FONT
            cell.alignment = _CENTER
            cell.border = _merged_row_border(col)
        row += 1
        for code in net.get("codes") or []:
            a_cell = lc_ws.cell(row=row, column=1)
            a_cell.border = _row_border(1)

            code_cell = lc_ws.cell(row=row, column=2, value=code["code_number"])
            code_cell.font = _CODE_FONT
            code_cell.alignment = _CENTER
            code_cell.border = _row_border(2)

            label_cell = lc_ws.cell(row=row, column=3, value=code["label"])
            label_cell.font = _LABEL_FONT
            label_cell.alignment = _LEFT_WRAP
            label_cell.border = _row_border(3)

            content_cell = lc_ws.cell(row=row, column=4)
            if code.get("content"):
                content_cell.value = code["content"]
            content_cell.font = _CONTENT_FONT
            content_cell.alignment = _LEFT_WRAP
            content_cell.border = _row_border(4)

            freq_cell = lc_ws.cell(row=row, column=5)
            pct_cell = lc_ws.cell(row=row, column=6)
            if code.get("frequency"):
                freq_cell.value = code["frequency"]
                pct_cell.value = code["pct"]
            freq_cell.font = _METRIC_FONT
            freq_cell.alignment = _CENTER
            freq_cell.border = _row_border(5)
            pct_cell.font = _METRIC_FONT
            pct_cell.alignment = _CENTER
            pct_cell.number_format = "0%"
            pct_cell.border = _row_border(6)

            row += 1

    last_row = row - 1
    if last_row >= 3:
        for col in range(1, _LC_LAST_COLUMN + 1):
            cell = lc_ws.cell(row=last_row, column=col)
            border = cell.border
            cell.border = Border(
                top=border.top, bottom=_MEDIUM, left=border.left, right=border.right
            )


def build_codebook_excel(export_data: dict) -> bytes:
    """Arma el .xlsx: una hoja LC_* por coding_type (Marcas y Razones tienen
    codebooks independientes, sección 2.7/MetaData.txt real) y UNA sola hoja
    Verbatims combinando ambos subproyectos por respondent_external_id — igual
    que el Case real de Ascribe (OnNextCase.txt), confirmado contra el fixture
    real de HBG Leches."""
    wb = Workbook()
    wb.remove(wb.active)

    study_name = export_data.get("study_name") or ""
    nets_by_coding_type = export_data.get("nets_by_coding_type") or {}
    for coding_type, nets in nets_by_coding_type.items():
        suffix = _sheet_suffix(coding_type)
        lc_ws = wb.create_sheet(f"LC {suffix}")
        _style_lc_sheet(lc_ws, f"{study_name} — {suffix}".strip(" —"))
        _write_lc_sheet(lc_ws, nets)

    header_fill = _HEADER_FILL
    header_font = Font(bold=True)

    verb_ws = wb.create_sheet("Verbatims")
    max_idx = max(export_data.get("max_output_index") or 0, 1)
    headers = ["Study ID", "Question ID", "Respondent ID", "Verbatim"]
    for i in range(1, max_idx + 1):
        headers.extend([f"Input ID {i}", f"Output ID {i}", f"Description {i}"])
    for col, header in enumerate(headers, start=1):
        cell = verb_ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    study_id = export_data.get("project_external_code") or ""
    for row_idx, verbatim in enumerate(export_data.get("verbatims") or [], start=2):
        verb_ws.cell(row=row_idx, column=1, value=study_id)
        verb_ws.cell(row=row_idx, column=2, value=verbatim.get("question_key") or "")
        verb_ws.cell(row=row_idx, column=3, value=verbatim.get("respondent_external_id") or "")
        verb_ws.cell(row=row_idx, column=4, value=verbatim.get("original_text") or "")
        codes_by_idx = {c["output_index"]: c for c in verbatim.get("codes", [])}
        for i in range(1, max_idx + 1):
            col_base = 5 + (i - 1) * 3
            code = codes_by_idx.get(i)
            if code:
                verb_ws.cell(row=row_idx, column=col_base, value=code["code_number"])
                verb_ws.cell(row=row_idx, column=col_base + 1, value=code["code_number"])
                verb_ws.cell(row=row_idx, column=col_base + 2, value=code["label"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
