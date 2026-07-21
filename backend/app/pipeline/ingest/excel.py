"""Parser del Excel real del cliente (Gloria Queso y similares)."""
import io
import random
import re

from openpyxl import load_workbook

_ALIAS_COUNT = re.compile(r"\((\d+)\)\s*$")


def sample_values(rows: list[dict[str, object]], column: str, n: int = 5) -> list[str]:
    """Hasta `n` valores no vacíos y distintos de `column`, elegidos al azar."""
    seen: list[str] = []
    for r in rows:
        v = r.get(column)
        if v is None or v == "":
            continue
        s = str(v).strip()
        if s and s not in seen:
            seen.append(s)
    return seen if len(seen) <= n else random.sample(seen, n)


def _row_values(ws, limit=8):
    return [[c for c in row] for row in ws.iter_rows(min_row=1, max_row=limit, values_only=True)]


def detect_sheets(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        head = _row_values(ws)
        flat = " ".join(str(c).lower() for row in head for c in row if c)
        if "verbatim" in flat and "respondent" in flat:
            guessed, cols = "verbatims", _header_of(head, "verbatim")
        elif "mencion principal" in flat or "mención principal" in flat or "codigo" in flat or "código" in flat:
            guessed, cols = "codebook", _header_of(head, "código", "codigo")
        else:
            guessed, cols = "unknown", []
        out.append({
            "name": ws.title,
            "rows": ws.max_row,
            "guessed_as": guessed,
            "columns": [str(c) for c in cols if c],
        })
    return out


def _header_of(head_rows, *keywords) -> list:
    for row in head_rows:
        low = [str(c).lower() if c else "" for c in row]
        if any(any(k in cell for k in keywords) for cell in low):
            return list(row)
    return []


def read_rows(content: bytes, sheet: str) -> tuple[list[str], list[dict[str, object]]]:
    """Primera fila como encabezado, resto como filas dict{columna: valor} —
    lectura genérica de una hoja para el preview de selección de campos
    (a diferencia de `parse_verbatims`, no asume una columna 'Verbatim')."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet]
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return [], []
    header = [str(c).strip() if c else "" for c in raw_rows[0]]
    columns = [c for c in header if c]
    rows = [dict(zip(header, r)) for r in raw_rows[1:]]
    return columns, rows


def parse_verbatims(content: bytes, sheet: str, column_map: dict[str, str] | None = None) -> list[dict]:
    """Lee verbatims desde Excel.

    La columna Respondent ID es obligatoria: es la clave del cruce con el perfil
    (REQ-05). Se puede indicar en `column_map['respondent_id']` o se detecta
    por alias ('Respondent ID', 'Serial', …). El .dau no usa esta función.
    """
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    cmap = column_map or {}
    resp_mapped = (cmap.get("respondent_id") or "").strip()
    verbatim_mapped = (cmap.get("verbatim") or "").strip()

    def _cells_lower(r) -> list[str]:
        return [str(c).strip().lower() for c in r if c is not None and str(c).strip()]

    def _is_header_row(r) -> bool:
        if not r:
            return False
        cells = _cells_lower(r)
        if verbatim_mapped and verbatim_mapped.lower() in cells:
            return True
        if resp_mapped and resp_mapped.lower() in cells:
            return True
        if "verbatim" in cells:
            return True
        if "respondent id" in cells or "serial" in cells:
            return True
        return False

    header_i = next((i for i, r in enumerate(rows) if _is_header_row(r)), None)
    if header_i is None:
        raise ValueError(
            f"La hoja '{sheet}' no tiene encabezado con Respondent ID / Verbatim"
        )
    header = [str(c).strip().lower() if c else "" for c in rows[header_i]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    # Respondent ID: mapeo explícito del wizard, o sniffing por alias.
    i_resp = None
    if resp_mapped and resp_mapped.lower() in header:
        i_resp = header.index(resp_mapped.lower())
    if i_resp is None:
        i_resp = col("respondent id", "serial", "respondent_id")
    if i_resp is None:
        raise ValueError(
            "El Excel debe tener una columna 'Respondent ID' "
            "(se valida contra el perfil de referencia). "
            "También acepta 'Serial'."
        )
    i_realid = col("realid", "real id")

    i_q = col("question id")
    # column_map.verbatim es explícito; si no viene, cae al sniffing actual.
    if verbatim_mapped and verbatim_mapped.lower() in header:
        i_txt = header.index(verbatim_mapped.lower())
    else:
        i_txt = col("verbatim")
    if i_txt is None:
        raise ValueError(f"La hoja '{sheet}' no tiene una columna 'Verbatim'")

    i_in = col("input id 1", "input id")
    i_out = col("output id 1", "output id")

    out = []
    for source_row_number, r in enumerate(rows[header_i + 1:], start=header_i + 2):
        if not r or not r[i_txt]:
            continue
        resp_val = r[i_resp] if i_resp < len(r) else None
        if resp_val is None or str(resp_val).strip() == "":
            continue
        out.append({
            "respondent_external_id": str(resp_val).strip(),
            "realid": str(r[i_realid]).strip() if i_realid is not None and i_realid < len(r) and r[i_realid] else None,
            "source_row_number": source_row_number,
            # Sin Question ID (Excel mínimo: Respondent ID + Verbatim) usamos
            # una clave estable; el wizard/ronda asigna Marcas o Razones.
            "question_key": (
                str(r[i_q]).strip()
                if i_q is not None and i_q < len(r) and r[i_q] is not None and str(r[i_q]).strip()
                else "verbatim"
            ),
            "original_text": str(r[i_txt]).strip(),
            "input_code_number": int(r[i_in]) if i_in is not None and i_in < len(r) and isinstance(r[i_in], (int, float)) else None,
            "truth_code_number": int(r[i_out]) if i_out is not None and i_out < len(r) and isinstance(r[i_out], (int, float)) else None,
        })
    return out


def parse_codebook(content: bytes, sheet: str) -> dict:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))

    codes, nets, current_net = [], [], None
    for r in rows:
        if not r or all(c is None for c in r):
            continue
        a, b, c, d = (r + (None,) * 4)[:4]

        if isinstance(a, str) and "neto" in a.lower():
            current_net = a.strip()
            if current_net not in nets:
                nets.append(current_net)
            continue

        if isinstance(b, int) and isinstance(c, str):
            codes.append(_mk_code(b, c, d, current_net))
            continue

        if isinstance(a, int) and isinstance(b, str):
            codes.append(_mk_code(a, b, c, current_net))
            continue

    return {"codes": codes, "nets": nets}


def _mk_code(number: int, label: str, contenido, net) -> dict:
    aliases = []
    if isinstance(contenido, str):
        for raw in contenido.split("/"):
            raw = raw.strip()
            if not raw:
                continue
            m = _ALIAS_COUNT.search(raw)
            occ = int(m.group(1)) if m else 1
            aliases.append({"text": _ALIAS_COUNT.sub("", raw).strip(), "occurrences": occ})
            
    label_str = str(label).strip()
    if "/" in label_str:
        parts = label_str.split("/", 1)
        if not net:
            net = parts[0].strip()
        label_str = parts[1].strip()
        
    formatted_label = label_str.capitalize() if label_str else label_str
    
    return {
        "code_number": number,
        "label": formatted_label,
        "aliases": aliases,
        "net": net,
        "is_system": 95 <= number <= 99,
    }
