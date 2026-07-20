"""Benchmark reproducible y de solo lectura para precisión de clasificación.

El benchmark cubre dos fuentes de evidencia independientes:

* inventario/reconciliación de un DAU contra un Excel de población esperada;
* replay temporal de libros históricos, aprendiendo únicamente de rondas previas.

No abre una sesión de base de datos ni modifica archivos de entrada. Los paths de
uploads se descubren por contenido y checksum, o se reciben explícitamente por CLI.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import load_workbook
from rapidfuzz import fuzz, process as rapidfuzz_process

from app.pipeline.ingest import dau, excel
from app.pipeline.normalize.text import normalize


CODING_TYPES = ("brands", "reasons")
HISTORICAL_SHEETS = {
    "brands": "Verbatims_Marcas",
    "reasons": "Verbatims_Razones",
}
EXPECTED_SHEETS = {"brands": "MARCAS", "reasons": "RAZONES"}
ROUND_RE = re.compile(r"(?:RONDA\s+|\bOL)(\d+)\b", re.IGNORECASE)
SHORT_TEXT_EXACT_MAX_LEN = 3
FUZZY_THRESHOLD = 80.0


@dataclass(frozen=True)
class LabeledVerbatim:
    coding_type: str
    original_text: str
    normalized_text: str
    code_number: int
    code_label: str | None = None


@dataclass(frozen=True)
class SourceVerbatim:
    coding_type: str
    respondent_external_id: str | None
    question_key: str | None
    original_text: str
    normalized_text: str


@dataclass(frozen=True)
class WorkbookSource:
    path: Path
    sha256: str
    round_number: int | None = None

    def public_identity(self) -> dict[str, object]:
        result: dict[str, object] = {
            "filename": self.path.name,
            "sha256": self.sha256,
        }
        if self.round_number is not None:
            result["round"] = self.round_number
        return result


class TemporalMemory:
    """Memoria exacta/fuzzy acumulada sin fuga de rondas futuras."""

    def __init__(self) -> None:
        self.counts: dict[str, dict[str, Counter[int]]] = {
            coding_type: {} for coding_type in CODING_TYPES
        }
        self.labels: dict[str, dict[int, Counter[str]]] = {
            coding_type: defaultdict(Counter) for coding_type in CODING_TYPES
        }

    def add(self, items: Iterable[LabeledVerbatim]) -> None:
        for item in items:
            by_text = self.counts[item.coding_type]
            if item.normalized_text not in by_text:
                by_text[item.normalized_text] = Counter()
            by_text[item.normalized_text][item.code_number] += 1
            if item.code_label:
                self.labels[item.coding_type][item.code_number][item.code_label] += 1

    def prediction(self, coding_type: str, normalized_text: str) -> int | None:
        counts = self.counts[coding_type].get(normalized_text)
        return _dominant(counts) if counts else None

    def conflicts(self) -> list[dict[str, object]]:
        result: list[dict[str, object]] = []
        for coding_type in CODING_TYPES:
            for normalized_text, counts in self.counts[coding_type].items():
                if len(counts) <= 1:
                    continue
                result.append(
                    {
                        "coding_type": coding_type,
                        "normalized_text": normalized_text,
                        "code_counts": {str(code): count for code, count in sorted(counts.items())},
                    }
                )
        return sorted(result, key=lambda row: (str(row["coding_type"]), str(row["normalized_text"])))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def discover_sources(root: Path) -> tuple[WorkbookSource, WorkbookSource, list[WorkbookSource]]:
    """Descubre DAU, población esperada y rondas históricas por contenido.

    Los duplicados físicos de uploads se colapsan por SHA-256. Si hay varios
    DAU válidos se elige el de mayor población; los callers que necesiten otro
    deben pasar rutas explícitas al CLI.
    """

    dau_candidates: list[tuple[int, WorkbookSource]] = []
    seen_dau: set[str] = set()
    for path in sorted(root.rglob("*.dau")):
        digest = sha256_file(path)
        if digest in seen_dau:
            continue
        seen_dau.add(digest)
        try:
            records = dau.parse_records(path.read_bytes())
        except dau.DauFormatError:
            continue
        dau_candidates.append((len(records), WorkbookSource(path, digest)))

    historical_by_checksum: dict[str, WorkbookSource] = {}
    expected_by_checksum: dict[str, WorkbookSource] = {}
    for path in sorted(root.rglob("*.xlsx")):
        digest = sha256_file(path)
        if digest in historical_by_checksum or digest in expected_by_checksum:
            continue
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_names = set(workbook.sheetnames)
            workbook.close()
        except Exception:
            continue

        if set(HISTORICAL_SHEETS.values()).issubset(sheet_names):
            round_number = infer_round_number(path)
            historical_by_checksum[digest] = WorkbookSource(path, digest, round_number)
        elif set(EXPECTED_SHEETS.values()).issubset(sheet_names):
            expected_by_checksum[digest] = WorkbookSource(path, digest)

    if not dau_candidates:
        raise ValueError(f"No se encontró un DAU válido debajo de {root}")
    if not expected_by_checksum:
        raise ValueError(f"No se encontró un Excel con hojas MARCAS/RAZONES debajo de {root}")

    dau_source = max(dau_candidates, key=lambda candidate: candidate[0])[1]
    expected_source = max(
        expected_by_checksum.values(),
        key=lambda source: source.path.stat().st_size,
    )
    historical = _one_workbook_per_round(historical_by_checksum.values())
    return dau_source, expected_source, historical


def infer_round_number(path: Path) -> int:
    match = ROUND_RE.search(path.name)
    if match:
        return int(match.group(1))

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in HISTORICAL_SHEETS.values():
            if sheet_name not in workbook.sheetnames:
                continue
            for row in workbook[sheet_name].iter_rows(min_row=1, max_row=5, values_only=True):
                for value in row:
                    if value is None:
                        continue
                    match = ROUND_RE.search(str(value))
                    if match:
                        return int(match.group(1))
    finally:
        workbook.close()
    raise ValueError(f"No se pudo inferir la ronda de {path.name}")


def _one_workbook_per_round(sources: Iterable[WorkbookSource]) -> list[WorkbookSource]:
    by_round: dict[int, WorkbookSource] = {}
    for source in sources:
        round_number = source.round_number
        if round_number is None:
            raise ValueError(f"Libro histórico sin ronda: {source.path}")
        previous = by_round.get(round_number)
        if previous and previous.sha256 != source.sha256:
            raise ValueError(
                f"Hay dos contenidos distintos para la ronda {round_number}: "
                f"{previous.path.name} y {source.path.name}"
            )
        by_round[round_number] = source
    rounds = sorted(by_round)
    if len(rounds) < 2:
        raise ValueError("Se necesitan al menos dos rondas históricas para el replay temporal")
    if rounds != list(range(rounds[0], rounds[-1] + 1)):
        raise ValueError(f"Las rondas históricas no son consecutivas: {rounds}")
    return [by_round[round_number] for round_number in rounds]


def read_historical_workbook(path: Path) -> dict[str, list[LabeledVerbatim]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            coding_type: _read_labeled_sheet(workbook[sheet_name], coding_type)
            for coding_type, sheet_name in HISTORICAL_SHEETS.items()
        }
    finally:
        workbook.close()


def _read_labeled_sheet(sheet, coding_type: str) -> list[LabeledVerbatim]:
    iterator = sheet.iter_rows(values_only=True)
    header = [str(value).strip().lower() if value is not None else "" for value in next(iterator)]
    text_index = _column_index(header, "verbatim")
    output_index = _column_index(header, "output id 1", "output id")
    input_index = _column_index(header, "input id 1", "input id")
    description_index = _column_index(header, "description 1", "description")
    if text_index is None or (output_index is None and input_index is None):
        raise ValueError(f"La hoja {sheet.title!r} no contiene Verbatim y un código de verdad")

    result: list[LabeledVerbatim] = []
    for row in iterator:
        if text_index >= len(row) or row[text_index] in (None, ""):
            continue
        raw_code = _first_present(row, output_index, input_index)
        code_number = _coerce_code(raw_code)
        if code_number is None:
            continue
        original_text = str(row[text_index]).strip()
        label = None
        if description_index is not None and description_index < len(row) and row[description_index] not in (None, ""):
            label = str(row[description_index]).strip()
        result.append(
            LabeledVerbatim(
                coding_type=coding_type,
                original_text=original_text,
                normalized_text=normalize(original_text),
                code_number=code_number,
                code_label=label,
            )
        )
    return result


def _column_index(header: Sequence[str], *names: str) -> int | None:
    for name in names:
        if name in header:
            return header.index(name)
    return None


def _first_present(row: Sequence[object], *indexes: int | None) -> object | None:
    for index in indexes:
        if index is not None and index < len(row) and row[index] not in (None, ""):
            return row[index]
    return None


def _coerce_code(value: object) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        number = float(str(value).strip())
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


def read_dau_source(path: Path) -> tuple[int, list[SourceVerbatim]]:
    content = path.read_bytes()
    records = dau.parse_records(content)
    rows: list[SourceVerbatim] = []
    for row in dau.parse_verbatims(content):
        coding_type = dau.guess_coding_type(str(row["question_key"]))
        if coding_type not in CODING_TYPES:
            continue
        original_text = str(row["original_text"])
        rows.append(
            SourceVerbatim(
                coding_type=coding_type,
                respondent_external_id=_optional_string(row.get("respondent_external_id")),
                question_key=_optional_string(row.get("question_key")),
                original_text=original_text,
                normalized_text=normalize(original_text),
            )
        )
    return len(records), rows


def read_expected_workbook(path: Path) -> list[SourceVerbatim]:
    content = path.read_bytes()
    result: list[SourceVerbatim] = []
    for coding_type, sheet_name in EXPECTED_SHEETS.items():
        for row in excel.parse_verbatims(content, sheet_name):
            original_text = str(row["original_text"])
            result.append(
                SourceVerbatim(
                    coding_type=coding_type,
                    respondent_external_id=_optional_string(row.get("respondent_external_id")),
                    question_key=_optional_string(row.get("question_key")),
                    original_text=original_text,
                    normalized_text=normalize(original_text),
                )
            )
    return result


def _optional_string(value: object) -> str | None:
    return None if value in (None, "") else str(value)


def build_inventory(
    dau_source: WorkbookSource,
    expected_source: WorkbookSource,
) -> dict[str, object]:
    record_count, dau_rows = read_dau_source(dau_source.path)
    expected_rows = read_expected_workbook(expected_source.path)
    by_type: dict[str, object] = {}
    for coding_type in CODING_TYPES:
        current_dau = [row for row in dau_rows if row.coding_type == coding_type]
        current_expected = [row for row in expected_rows if row.coding_type == coding_type]
        by_type[coding_type] = {
            "dau": _group_distribution(current_dau),
            "expected": _group_distribution(current_expected),
            "reconciliation": _reconcile(current_dau, current_expected),
        }

    return {
        "sources": {
            "dau": dau_source.public_identity(),
            "expected_workbook": expected_source.public_identity(),
        },
        "dau_records": record_count,
        "dau_verbatims": len(dau_rows),
        "expected_verbatims": len(expected_rows),
        "question_variables": len({row.question_key for row in dau_rows if row.question_key}),
        "by_coding_type": by_type,
    }


def _group_distribution(rows: Sequence[SourceVerbatim]) -> dict[str, object]:
    frequencies = Counter(row.normalized_text for row in rows)
    top_twenty = sum(count for _, count in frequencies.most_common(20))
    return {
        "rows": len(rows),
        "groups": len(frequencies),
        "singletons": sum(count == 1 for count in frequencies.values()),
        "top_20_row_share": _ratio(top_twenty, len(rows)),
    }


def _reconcile(
    dau_rows: Sequence[SourceVerbatim],
    expected_rows: Sequence[SourceVerbatim],
) -> dict[str, object]:
    dau_text = Counter(row.normalized_text for row in dau_rows)
    expected_text = Counter(row.normalized_text for row in expected_rows)
    dau_identity = Counter(
        (row.respondent_external_id, row.normalized_text) for row in dau_rows
    )
    expected_identity = Counter(
        (row.respondent_external_id, row.normalized_text) for row in expected_rows
    )
    return {
        "expected_is_text_multiset_subset": not bool(expected_text - dau_text),
        "expected_is_respondent_text_multiset_subset": not bool(expected_identity - dau_identity),
        "dau_extra_rows_by_text": sum((dau_text - expected_text).values()),
        "expected_missing_rows_by_text": sum((expected_text - dau_text).values()),
        "dau_extra_rows_by_respondent_text": sum((dau_identity - expected_identity).values()),
        "expected_missing_rows_by_respondent_text": sum((expected_identity - dau_identity).values()),
    }


def temporal_replay(rounds: Mapping[int, Mapping[str, Sequence[LabeledVerbatim]]]) -> dict[str, object]:
    if len(rounds) < 2:
        raise ValueError("El replay temporal necesita al menos dos rondas")

    memory = TemporalMemory()
    replay_rows: list[dict[str, object]] = []
    first_round = min(rounds)
    for round_number in sorted(rounds):
        if round_number != first_round:
            for coding_type in CODING_TYPES:
                current = list(rounds[round_number].get(coding_type, []))
                exact_predictions = {
                    item.normalized_text: memory.prediction(coding_type, item.normalized_text)
                    for item in current
                }
                exact_metrics = evaluate_predictions(current, exact_predictions)
                fuzzy_predictions, fuzzy_details = fuzzy_candidates(current, memory, coding_type)
                uncovered = [
                    item for item in current if exact_predictions[item.normalized_text] is None
                ]
                fuzzy_metrics = evaluate_predictions(uncovered, fuzzy_predictions)
                combined_predictions = dict(exact_predictions)
                combined_predictions.update(
                    {
                        text: prediction
                        for text, prediction in fuzzy_predictions.items()
                        if prediction is not None
                    }
                )
                replay_rows.append(
                    {
                        "round": round_number,
                        "coding_type": coding_type,
                        "exact": exact_metrics,
                        "fuzzy_on_exact_residue": fuzzy_metrics,
                        "exact_plus_fuzzy": evaluate_predictions(current, combined_predictions),
                        "fuzzy_false_positive_examples": fuzzy_details,
                    }
                )
        for coding_type in CODING_TYPES:
            memory.add(rounds[round_number].get(coding_type, []))

    return {
        "first_training_round": first_round,
        "evaluations": replay_rows,
        "historical_conflicts": memory.conflicts(),
    }


def fuzzy_candidates(
    current: Sequence[LabeledVerbatim],
    memory: TemporalMemory,
    coding_type: str,
) -> tuple[dict[str, int | None], list[dict[str, object]]]:
    """Genera candidatos WRatio sólo para el residuo no cubierto por exact."""

    aliases = list(memory.counts[coding_type])
    predictions: dict[str, int | None] = {}
    examples: list[dict[str, object]] = []
    actual_by_text = _actual_counts(current)
    for normalized_text in dict.fromkeys(item.normalized_text for item in current):
        if memory.prediction(coding_type, normalized_text) is not None:
            continue
        match = rapidfuzz_process.extractOne(
            normalized_text,
            aliases,
            scorer=fuzz.WRatio,
            score_cutoff=FUZZY_THRESHOLD,
        ) if aliases else None
        if not match or (
            len(normalized_text) <= SHORT_TEXT_EXACT_MAX_LEN and float(match[1]) < 100.0
        ):
            predictions[normalized_text] = None
            continue
        matched_alias, score, _ = match
        predicted = memory.prediction(coding_type, matched_alias)
        predictions[normalized_text] = predicted
        actual = _dominant(actual_by_text[normalized_text])
        if predicted != actual and len(examples) < 10:
            examples.append(
                {
                    "text": normalized_text,
                    "matched_alias": matched_alias,
                    "score": round(float(score), 4),
                    "predicted_code": predicted,
                    "actual_code": actual,
                }
            )
    return predictions, examples


def evaluate_predictions(
    current: Sequence[LabeledVerbatim],
    predictions: Mapping[str, int | None],
) -> dict[str, object]:
    row_confusion: Counter[tuple[str, str]] = Counter()
    row_covered = row_correct = 0
    for item in current:
        prediction = predictions.get(item.normalized_text)
        predicted_label = str(prediction) if prediction is not None else "__abstain__"
        row_confusion[(predicted_label, str(item.code_number))] += 1
        if prediction is not None:
            row_covered += 1
            row_correct += prediction == item.code_number

    actual_by_text = _actual_counts(current)
    group_confusion: Counter[tuple[str, str]] = Counter()
    group_covered = group_correct = 0
    ambiguous_groups = 0
    for normalized_text, actual_counts in actual_by_text.items():
        prediction = predictions.get(normalized_text)
        actual = _dominant(actual_counts)
        ambiguous_groups += len(actual_counts) > 1
        predicted_label = str(prediction) if prediction is not None else "__abstain__"
        group_confusion[(predicted_label, str(actual))] += 1
        if prediction is not None:
            group_covered += 1
            group_correct += prediction == actual

    return {
        "row_weighted": _metric_block(
            total=len(current),
            covered=row_covered,
            correct=row_correct,
            confusion=row_confusion,
        ),
        "group_weighted": {
            **_metric_block(
                total=len(actual_by_text),
                covered=group_covered,
                correct=group_correct,
                confusion=group_confusion,
            ),
            "ambiguous_truth_groups": ambiguous_groups,
        },
    }


def _actual_counts(items: Sequence[LabeledVerbatim]) -> dict[str, Counter[int]]:
    result: dict[str, Counter[int]] = defaultdict(Counter)
    for item in items:
        result[item.normalized_text][item.code_number] += 1
    return result


def _dominant(counts: Counter[int]) -> int:
    # Counter preserva el orden de primera observación para empates. Esa regla
    # es estable y replica cómo se acumuló la memoria histórica original.
    return counts.most_common(1)[0][0]


def _metric_block(
    *,
    total: int,
    covered: int,
    correct: int,
    confusion: Counter[tuple[str, str]],
) -> dict[str, object]:
    return {
        "total": total,
        "covered": covered,
        "correct": correct,
        "coverage": _ratio(covered, total),
        "precision_on_covered": _ratio(correct, covered),
        "recall_with_abstention": _ratio(correct, total),
        "abstention_rate": _ratio(total - covered, total),
        "confusion_matrix": {
            predicted: {
                actual: confusion[(predicted, actual)]
                for actual in sorted({key[1] for key in confusion if key[0] == predicted})
            }
            for predicted in sorted({key[0] for key in confusion})
        },
    }


def _ratio(numerator: int, denominator: int) -> float | None:
    return round(numerator / denominator, 8) if denominator else None


def build_report(
    dau_source: WorkbookSource,
    expected_source: WorkbookSource,
    historical_sources: Sequence[WorkbookSource],
) -> dict[str, object]:
    rounds = {
        source.round_number: read_historical_workbook(source.path)
        for source in historical_sources
        if source.round_number is not None
    }
    return {
        "benchmark_version": "classification-accuracy-v1",
        "normalization": "app.pipeline.normalize.text.normalize",
        "fuzzy_policy": {
            "algorithm": "rapidfuzz.fuzz.WRatio",
            "threshold": FUZZY_THRESHOLD,
            "short_text_exact_max_length": SHORT_TEXT_EXACT_MAX_LEN,
            "purpose": "candidate_evaluation_only",
        },
        "inventory": build_inventory(dau_source, expected_source),
        "historical_sources": [source.public_identity() for source in historical_sources],
        "temporal_replay": temporal_replay(rounds),
    }


def render_markdown(report: Mapping[str, object]) -> str:
    inventory = report["inventory"]
    assert isinstance(inventory, Mapping)
    by_type = inventory["by_coding_type"]
    assert isinstance(by_type, Mapping)
    replay = report["temporal_replay"]
    assert isinstance(replay, Mapping)
    evaluations = replay["evaluations"]
    assert isinstance(evaluations, list)

    lines = [
        "# Baseline reproducible de precisión de clasificación",
        "",
        "Este reporte es generado por `scripts.classification_accuracy_benchmark` sin escribir en la base de datos.",
        "",
        "## Inventario y reconciliación",
        "",
        f"- Registros DAU: {inventory['dau_records']}",
        f"- Verbatims DAU: {inventory['dau_verbatims']}",
        f"- Verbatims Excel esperado: {inventory['expected_verbatims']}",
        f"- Variables de pregunta DAU: {inventory['question_variables']}",
        "",
        "| Tipo | Filas DAU | Grupos DAU | Singletons DAU | Top 20 | Filas esperadas | Grupos esperados | Extra DAU |",
        "|---|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for coding_type in CODING_TYPES:
        section = by_type[coding_type]
        dau_section = section["dau"]
        expected_section = section["expected"]
        reconciliation = section["reconciliation"]
        lines.append(
            f"| {coding_type} | {dau_section['rows']} | {dau_section['groups']} | "
            f"{dau_section['singletons']} | {_percent(dau_section['top_20_row_share'])} | "
            f"{expected_section['rows']} | {expected_section['groups']} | "
            f"{reconciliation['dau_extra_rows_by_respondent_text']} |"
        )

    lines.extend(
        [
            "",
            "## Replay temporal",
            "",
            "Cada ronda aprende exclusivamente de rondas anteriores. La precisión fuzzy se mide sólo sobre el residuo no cubierto por memoria exacta.",
            "",
            "| Tipo | Ronda | Cobertura exacta (filas) | Precisión exacta | Candidatos fuzzy | Precisión fuzzy |",
            "|---|---:|---:|---:|---:|---:|",
        ]
    )
    for evaluation in evaluations:
        exact_rows = evaluation["exact"]["row_weighted"]
        fuzzy_rows = evaluation["fuzzy_on_exact_residue"]["row_weighted"]
        lines.append(
            f"| {evaluation['coding_type']} | R{evaluation['round']} | "
            f"{_percent(exact_rows['coverage'])} | {_percent(exact_rows['precision_on_covered'])} | "
            f"{fuzzy_rows['covered']} | {_percent(fuzzy_rows['precision_on_covered'])} |"
        )

    conflicts = replay["historical_conflicts"]
    lines.extend(
        [
            "",
            "## Conflictos históricos",
            "",
            f"Se detectaron {len(conflicts)} textos normalizados con más de un código histórico.",
            "",
            "| Tipo | Texto normalizado | Conteos por código |",
            "|---|---|---|",
        ]
    )
    for conflict in conflicts:
        counts = ", ".join(f"{code}: {count}" for code, count in conflict["code_counts"].items())
        lines.append(f"| {conflict['coding_type']} | {conflict['normalized_text']} | {counts} |")

    lines.extend(
        [
            "",
            "## Reproducción",
            "",
            "```bash",
            "cd backend",
            "PYTHONPATH=. env/bin/python -m scripts.classification_accuracy_benchmark --workspace-root .. \\",
            "  --json-output ../tasks/reports/classification_accuracy_baseline.json \\",
            "  --markdown-output ../tasks/reports/classification_accuracy_baseline.md",
            "```",
            "",
        ]
    )
    return "\n".join(lines)


def _percent(value: object) -> str:
    if value is None:
        return "n/a"
    return f"{float(value) * 100:.2f}%"


def _source_from_explicit(path: Path, round_number: int | None = None) -> WorkbookSource:
    if not path.is_file():
        raise ValueError(f"No existe el archivo {path}")
    return WorkbookSource(path=path, sha256=sha256_file(path), round_number=round_number)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace-root", type=Path, default=Path(".."))
    parser.add_argument("--dau", type=Path)
    parser.add_argument("--expected-xlsx", type=Path)
    parser.add_argument(
        "--historical-xlsx",
        type=Path,
        action="append",
        help="Libro histórico; repetir una vez por ronda. La ronda se infiere del nombre/contenido.",
    )
    parser.add_argument("--json-output", type=Path)
    parser.add_argument("--markdown-output", type=Path)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    explicit = args.dau or args.expected_xlsx or args.historical_xlsx
    if explicit and not (args.dau and args.expected_xlsx and args.historical_xlsx):
        raise SystemExit(
            "Si se proporciona una fuente explícita, se requieren --dau, "
            "--expected-xlsx y al menos dos --historical-xlsx"
        )

    if explicit:
        dau_source = _source_from_explicit(args.dau)
        expected_source = _source_from_explicit(args.expected_xlsx)
        historical_sources = _one_workbook_per_round(
            _source_from_explicit(path, infer_round_number(path))
            for path in args.historical_xlsx
        )
    else:
        dau_source, expected_source, historical_sources = discover_sources(
            args.workspace_root.resolve()
        )

    report = build_report(dau_source, expected_source, historical_sources)
    json_text = json.dumps(report, ensure_ascii=False, indent=2, sort_keys=False) + "\n"
    markdown_text = render_markdown(report)

    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(json_text, encoding="utf-8")
    if args.markdown_output:
        args.markdown_output.parent.mkdir(parents=True, exist_ok=True)
        args.markdown_output.write_text(markdown_text, encoding="utf-8")
    if not args.json_output and not args.markdown_output:
        print(json_text, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
